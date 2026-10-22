"""Поиск артикулов в базе и заполнение Excel-выгрузки Битрикса.

Правила (CLAUDE.md): значения только из каталога или артикула; заполняются только пустые ячейки;
в листах загрузки нет пустых ячеек; код товара — текст из 8 знаков.
"""
import re
import sqlite3
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from articles import decode, norm  # noqa: E402

YELLOW = PatternFill("solid", fgColor="FFF2A8")
BLUE = PatternFill("solid", fgColor="D9E8FB")
HEAD = PatternFill("solid", fgColor="DDE6F0")

# Какие каталоги смотреть для бренда из файла (порядок = приоритет)
BRAND_CATALOGS = {
    "metalcraft": ["UNKNOWN112", "YOUYIQIE", "WEIDING"],
    "kkdaoku": ["UNKNOWN112", "YOUYIQIE"],
    "youyiqie": ["YOUYIQIE", "UNKNOWN112"],
    "sant": ["SANTTURN", "SANT2022"],
}
CATALOG_TITLE = {}


def fmt_num(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else round(float(v), 3)


# ---------- шаблоны колонок сайта ----------
def _bore(p, a, dec):
    d = p.get("d")
    out = {
        "d, мм": d, "S, мм": p.get("f"), "e, мм": p.get("L1"), "Dmin, мм": p.get("Dmin"),
        "Мин. диаметр растачивания, мм": p.get("Dmin"), "H, мм": p.get("H"), "L, мм": p.get("L"),
        "Длина резца, мм": p.get("L"), "Диаметр хвостовика, мм": d,
        "Сечение резца, мм": f"Ø{fmt_num(d)}" if d else None,
    }
    return out


def _carbide(p, a, dec):
    d, A = p.get("d"), p.get("A")
    return {
        "d, мм": d, "S, мм": (A - d / 2) if (A is not None and d) else None,
        "Dmin, мм": p.get("Dmin"), "Мин. диаметр растачивания, мм": p.get("Dmin"),
        "L, мм": p.get("L"), "Длина резца, мм": p.get("L"), "Диаметр хвостовика, мм": d,
        "Сечение резца, мм": f"Ø{fmt_num(d)}" if d else None,
    }


def _thread(p, a, dec):
    d, size, hand = p.get("d"), dec.get("insert_size"), dec.get("hand")
    return {
        "L, мм": p.get("L"), "Длина резца, мм": p.get("L"), "Dmin, мм": p.get("Dmin"),
        "Мин. диаметр растачивания, мм": p.get("Dmin"), "d, мм": d, "Диаметр хвостовика, мм": d, "b, мм": d,
        "Сечение резца, мм": f"Ø{fmt_num(d)}" if d else None, "h, мм": p.get("H"), "H, мм": p.get("H"),
        "S, мм": p.get("F"), "L1, мм": p.get("L1"),
        "Размер пластины": int(size) if size else None,
        "Устанавливаемая пластина": f"{size}I{hand}.." if size and hand else None,
        # Константы, принятые на сайте для резьбовых расточных державок
        "Крепление пластины": "S - Фиксация винтом ", "Вид точения": "Внутреннее", "Тип державки": "S (стальная)",
        "Направление резьбы": "Правое" if hand == "R" else "Внутренняя",
    }


TEMPLATES = {"holder_bore": _bore, "holder_bore_carbide": _carbide, "holder_thread": _thread}


@dataclass
class Match:
    article: str
    catalog: str
    group: str
    page_pdf: int | None
    params: dict
    notes: dict = field(default_factory=dict)

    @property
    def source(self):
        t = CATALOG_TITLE.get(self.catalog, self.catalog)
        return f"{t}, стр. PDF {self.page_pdf}" if self.page_pdf else t


class Engine:
    def __init__(self, db_path):
        self.con = sqlite3.connect(db_path)
        self.con.row_factory = sqlite3.Row
        for r in self.con.execute("SELECT code, title FROM catalog"):
            CATALOG_TITLE[r["code"]] = r["title"]

    # ---------- поиск ----------
    def find(self, article: str, brand: str | None = None) -> list[Match]:
        rows = self.con.execute(
            """SELECT i.id, i.article_raw, c.code cat, g.code grp, p.pdf_page
               FROM item i JOIN series s ON s.id=i.series_id JOIN catalog c ON c.id=s.catalog_id
               JOIN product_group g ON g.id=s.group_id LEFT JOIN page p ON p.id=i.page_id
               WHERE i.article_norm=?""", (norm(article),)).fetchall()
        out = []
        for r in rows:
            params, notes = {}, {}
            for pr in self.con.execute("SELECT param_code, value_num, note FROM item_param WHERE item_id=?", (r["id"],)):
                params[pr["param_code"]] = pr["value_num"]
                notes[pr["param_code"]] = pr["note"]
            out.append(Match(r["article_raw"], r["cat"], r["grp"], r["pdf_page"], params, notes))
        order = BRAND_CATALOGS.get((brand or "").strip().lower())
        if order is not None:
            out = [m for m in out if m.catalog in order]
            out.sort(key=lambda m: order.index(m.catalog))
        return out

    def series_hint(self, article: str) -> list[str]:
        """Где в каталогах встречается серия (для «нет в базе»)."""
        dec = decode(article)
        s = dec.get("series")
        if not s:
            return []
        res = self.con.execute(
            """SELECT c.title, p.pdf_page FROM page p JOIN catalog c ON c.id=p.catalog_id
               WHERE upper(p.ocr_text) LIKE ? LIMIT 8""", (f"%{s}%",)).fetchall()
        return [f"{r['title']}, стр. PDF {r['pdf_page']}" for r in res]

    def site_values(self, article: str, brand: str | None = None) -> tuple[dict, str, Match | None]:
        """Значения для колонок сайта: из каталога, иначе только из артикула."""
        dec = decode(article)
        ms = self.find(article, brand)
        hand_txt = {"R": "R (правое)", "L": "L (левое)"}.get(dec.get("hand"))
        if ms:
            m = ms[0]
            vals = {k: v for k, v in TEMPLATES[m.group](m.params, article, dec).items() if v is not None}
            if hand_txt:
                vals["Исполнение"] = hand_txt
            src = {k: m.source for k in vals}
            if m.group == "holder_bore_carbide" and "S, мм" in vals:
                src["S, мм"] = m.source + " (S = A − d/2)"
            return vals, src, m
        vals, src = {}, {}
        if hand_txt:
            vals["Исполнение"], src["Исполнение"] = hand_txt, "артикул"
        if dec.get("d"):
            vals["Сечение резца, мм"], src["Сечение резца, мм"] = f"Ø{dec['d']}", "артикул"
        if dec.get("L"):
            vals["Длина резца, мм"], src["Длина резца, мм"] = dec["L"], "артикул (буква длины)"
        return vals, src, None


# ---------- Excel ----------
def _empty(v):
    return v is None or str(v).strip() in ("", "-")


def _num(v):
    try:
        return float(str(v).strip().replace(",", ".").replace("Ø", ""))
    except (TypeError, ValueError):
        return None


def _code(v):
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.zfill(8) if s.isdigit() else s


def read_table(path):
    """Читает выгрузку: строка заголовков — первая, где есть «Артикул»."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hdr_row = None
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if "Артикул" in vals:
            hdr_row, headers = r, vals
            break
    if hdr_row is None:
        raise ValueError("В файле нет колонки «Артикул» в первых 10 строках")
    rows = []
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        d = {h: v for h, v in zip(headers, r) if h}
        a = d.get("Артикул")
        if a and str(a).strip() and str(a).strip() != "Итог":
            rows.append(d)
    return headers, rows


@dataclass
class FillResult:
    rows: list
    found: int
    filled_cells: int
    conflicts: list
    not_found: list
    sources: list


def fill_rows(engine: Engine, headers, rows, seo_sort_by: str | None = None) -> FillResult:
    out, conflicts, not_found, sources = [], [], [], []
    found = filled = 0
    for r in rows:
        art = str(r["Артикул"]).strip()
        brand = r.get("Бренд")
        vals, src, m = engine.site_values(art, brand)
        dec = decode(art)
        if m:
            found += 1
            # защита: d на сайте не совпадает с каталогом → строку не трогаем
            site_d = _num(r.get("d, мм")) or _num(r.get("Сечение резца, мм"))
            if site_d and m.params.get("d") and abs(site_d - m.params["d"]) > 0.001:
                not_found.append([art, _code(r.get("_Код товара", "")), "d на сайте "
                                  f"{fmt_num(site_d)}, в каталоге {fmt_num(m.params['d'])} — строку не заполнял, проверьте"])
                continue
            # буква длины в артикуле против L каталога
            if dec.get("L") and m.params.get("L") and dec["L"] != m.params["L"]:
                for k in ("L, мм", "Длина резца, мм"):
                    vals.pop(k, None)
                conflicts.append([art, _code(r.get("_Код товара", "")), "L", f"буква длины = {dec['L']}",
                                  fmt_num(m.params["L"]), m.source])
        final, fl = {}, set()
        # Колонки свойств из шаблона добавляем, даже если их нет в файле (голый список «артикул + код»).
        eff_headers = list(headers) + [c for c in sorted(TEMPLATE_COLUMNS) if c and c not in headers]
        for h in eff_headers:
            if not h:
                continue
            cur = r.get(h)
            if h in vals:
                v = vals[h]
                v = fmt_num(v) if isinstance(v, float) else v
                if _empty(cur):
                    final[h] = v
                    fl.add(h)
                    sources.append([art, _code(r.get("_Код товара", "")), h, v, src.get(h, "")])
                else:
                    final[h] = cur
                    if _num(cur) is not None and _num(v) is not None and abs(_num(cur) - _num(v)) > 0.001:
                        conflicts.append([art, _code(r.get("_Код товара", "")), h, cur, v, src.get(h, "")])
            elif not _empty(cur):
                final[h] = cur
        missing = [h for h in headers if h in TEMPLATE_COLUMNS and _empty(final.get(h))]
        if missing:
            why = "нет в базе каталогов" if not m else "в каталоге нет этих колонок"
            hint = "" if m else ("; серия встречается: " + "; ".join(engine.series_hint(art)[:3])) if engine.series_hint(art) else ""
            not_found.append([art, _code(r.get("_Код товара", "")), f"пусто: {', '.join(missing)} — {why}{hint}"])
        filled += len(fl)
        final["Артикул"] = art
        final["_Код товара"] = _code(r.get("_Код товара", ""))
        out.append({"values": final, "filled": fl, "changed": set()})
    if seo_sort_by:
        out.sort(key=lambda x: (_num(x["values"].get(seo_sort_by)) or 0,
                                _num(x["values"].get("Dmin, мм")) or 0, x["values"]["Артикул"]))
        for i, x in enumerate(out, 1):
            if str(x["values"].get("_SEO_Sort")) != str(i * 10):
                x["changed"].add("_SEO_Sort")
            x["values"]["_SEO_Sort"] = i * 10
    return FillResult(out, found, filled, conflicts, not_found, sources)


TEMPLATE_COLUMNS = set()
for _f in TEMPLATES.values():
    TEMPLATE_COLUMNS |= set(_f({"d": 1, "A": 1, "H": 1, "L": 1, "L1": 1, "f": 1, "F": 1, "Dmin": 1}, "", {"insert_size": "16", "hand": "R"}).keys())
TEMPLATE_COLUMNS |= {"Исполнение"}


def write_result(res: FillResult, headers, path, seo_sort_by=None):
    """Листы загрузки без пустых ячеек: строки группируются по набору заполненных колонок."""
    wanted = [h for h in headers if h and (h in TEMPLATE_COLUMNS or h in ("_Код товара", "Артикул"))]
    # Колонки, заполненные из шаблона, но отсутствовавшие в исходном файле (голый список).
    for c in sorted({h for x in res.rows for h in (x["filled"] | x["changed"])}):
        if c and c not in wanted:
            wanted.append(c)
    if seo_sort_by and "_SEO_Sort" not in wanted:
        wanted.append("_SEO_Sort")
    groups = {}
    for x in res.rows:
        if not x["filled"] and not x["changed"]:
            continue
        cols = tuple(h for h in wanted if not _empty(x["values"].get(h)))
        groups.setdefault(cols, []).append(x)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for n, (cols, items) in enumerate(sorted(groups.items(), key=lambda kv: -len(kv[1])), 1):
        cols = ["Артикул", "_Код товара"] + [c for c in cols if c not in ("Артикул", "_Код товара")]
        ws = wb.create_sheet("Загрузка" if n == 1 else f"Загрузка {n}")
        ws.append(cols)
        for x in items:
            ws.append([x["values"].get(c) for c in cols])
            rr = ws.max_row
            ws.cell(rr, 2).number_format = "@"
            for j, c in enumerate(cols, 1):
                if c in x["filled"]:
                    ws.cell(rr, j).fill = YELLOW
                elif c in x["changed"]:
                    ws.cell(rr, j).fill = BLUE
        _style(ws, cols)
    if not wb.sheetnames:
        wb.create_sheet("Загрузка").append(["Нечего заполнять: пустых ячеек, которые есть в каталогах, нет"])
    for title, head, data in (("Источники", ["Артикул", "Код", "Колонка", "Значение", "Источник"], res.sources),
                              ("Расходится с каталогом", ["Артикул", "Код", "Колонка", "На сайте", "В каталоге", "Источник"], res.conflicts),
                              ("Не найдено", ["Артикул", "Код", "Причина"], res.not_found)):
        ws = wb.create_sheet(title)
        ws.append(head)
        for row in data:
            ws.append(row)
            ws.cell(ws.max_row, 2).number_format = "@"
        _style(ws, head, wide=True)
    wb.save(path)


def _style(ws, cols, wide=False):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 16 if i <= 2 else (60 if wide and i >= 3 else 12)
    ws.freeze_panes = "B2"


def write_cards(headers, rows, path):
    """Прямой экспорт карточек, взятых с cnc1.ru, в Excel-выгрузку (без заполнения из каталога).

    Один лист «Выгрузка с сайта»: Артикул, _Код товара (текст, 8 знаков), затем все свойства карточек
    в том порядке, как их отдал сайт. Пустые значения («-») остаются пустыми. Лист «Ссылки» — URL карточек.
    """
    cols = ["Артикул", "_Код товара"] + [h for h in headers if h not in ("Артикул", "_Код товара", "_URL")]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Выгрузка с сайта")
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
        ws.cell(ws.max_row, 2).number_format = "@"   # код — текстом, ведущие нули сохраняются
    _style(ws, cols)
    if any(r.get("_URL") for r in rows):
        ws2 = wb.create_sheet("Ссылки")
        ws2.append(["Артикул", "_Код товара", "URL карточки"])
        for r in rows:
            ws2.append([r.get("Артикул"), r.get("_Код товара"), r.get("_URL")])
            ws2.cell(ws2.max_row, 2).number_format = "@"
        _style(ws2, ["Артикул", "_Код товара", "URL карточки"], wide=True)
    wb.save(path)


def process_file(engine, in_path, out_path, seo_sort_by=None) -> FillResult:
    headers, rows = read_table(in_path)
    res = fill_rows(engine, headers, rows, seo_sort_by)
    write_result(res, headers, out_path, seo_sort_by)
    return res
