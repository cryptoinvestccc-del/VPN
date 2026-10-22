import re, sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
D = '/tmp/claude-0/-home-claude/77365e1c-a1c8-5008-b397-5c19b94b801a/scratchpad/xl/'
sys.path.insert(0, D)
from kk import KK
import fill as F  # reuse PT (PTHK) and SANT tables; fill.py runs on import, harmless

C_D, C_S, C_E, C_DM, C_H, C_L, C_DM2, C_SEC, C_LR = F.CORE
CORE = F.CORE
ISO = dict(F=80, G=90, H=100, J=110, K=125, L=140, M=150, N=160, P=170, Q=180, R=200, S=250, T=300, U=350, V=400)
E = lambda v: v in (None, '', '-', 0, '0')
def num(v):
    m = re.search(r'-?\d+(?:[.,]\d+)?', str(v)) if v is not None else None
    return float(m.group().replace(',', '.')) if m else None
def nice(x): return int(x) if isinstance(x, float) and x == int(x) else x

wb = openpyxl.load_workbook(D + 'src.xlsx'); ws = wb.active
hdr = [c.value for c in ws[1]]; col = {h: i + 1 for i, h in enumerate(hdr)}
Y = PatternFill('solid', fgColor='FFF2A8')   # было пусто
B = PatternFill('solid', fgColor='BDD7EE')   # исправлено
O = PatternFill('solid', fgColor='FFD8A8')   # по аналогу
filled, fixed, notdone, kept = [], [], [], []

for r in range(2, ws.max_row + 1):
    row = {h: ws.cell(r, col[h]).value for h in hdr}
    art = row['Артикул'] or ''; brand = row['Бренд']
    vals, src, how, keepL = {}, '', 'из каталога', False
    m = re.match(r'^([A-Z])(\d{2})([A-Z])-([A-Z]{4})([RL])(\d{2})(H?)$', art)
    if art in F.SANT:
        vals, src, how = F.SANT[art]
    elif m and brand in ('MetalCraft', 'kkdaoku'):
        t, dd, let, ser, hand, sz, h = m.groups(); pre = t + dd + let
        if t == 'S':
            hit = KK.get((pre, ser + sz + h)) or (KK.get((pre, ser[:3] + '_' + sz)) if ser[:2] == 'SV' else None)
            if hit:
                (dm, d, L, L1, H, f), page = hit
                vals = {C_D: d, C_S: f, C_E: L1, C_DM: dm, C_H: H, C_L: L, C_DM2: dm, C_SEC: f'Ø{d}', C_LR: L}
                src = f'каталог MetalCraft, стр. {page} PDF ({pre}-{ser}R/L{sz}{h})'
                if H >= d:
                    del vals[C_H]
                    kept.append([art, brand, f'H в каталоге {H} больше диаметра {d} — опечатка каталога, H не трогал', src])
                if ISO.get(let) and L != ISO[let]:
                    keepL = True
        elif t in 'CEH' and brand == 'MetalCraft':
            hit = F.PT.get((pre, ser + sz))
            if hit:
                d, L, A, dm, page = hit
                vals = {C_D: d, C_S: nice(A - d / 2), C_DM: dm, C_L: L, C_DM2: dm, C_SEC: f'Ø{d}', C_LR: L}
                src, how = f'PTHK/Weiding, стр. {page} каталога', 'из каталога (S = A − d/2)'
    if vals and not E(row[C_D]) and C_D in vals and num(row[C_D]) != num(vals[C_D]):
        kept.append([art, brand, f'd на сайте {row[C_D]}, в каталоге {vals[C_D]} — строку не трогал, проверьте диаметр', src])
        vals = {}
    f_, x_ = [], []
    for k, v in vals.items():
        if keepL and k in (C_L, C_LR):
            if num(row[k]) != num(v):
                kept.append([art, brand, f'{k.split(",")[0]}: на сайте {row[k]}, в каталоге {v}, но буква длины в артикуле = {ISO[let]} мм — оставил как на сайте', src])
            continue
        c = ws.cell(r, col[k]); cur = c.value
        if E(cur):
            c.value = v; c.fill = O if how == 'по аналогу' else Y; f_.append(f'{k.split(",")[0]}={v}')
        elif (num(cur) is not None and num(v) is not None and abs(num(cur) - num(v)) > 0.001) :
            c.value = v; c.fill = B; x_.append(f'{k.split(",")[0]}: {cur} → {v}')
    for a, b in ((C_DM, C_DM2), (C_DM2, C_DM)):
        va, cb = ws.cell(r, col[a]).value, ws.cell(r, col[b])
        if E(cb.value) and not E(va):
            cb.value = va; cb.fill = Y; f_.append(f'{b.split(",")[0]}={va} (дубль {a.split(",")[0]})')
            src = src or 'та же строка'
    if f_: filled.append([art, brand, ', '.join(f_), how, src])
    if x_: fixed.append([art, brand, '; '.join(x_), src])
    still = [k.split(',')[0] for k in CORE if E(ws.cell(r, col[k]).value)]
    if still:
        if vals and m and m.group(1) in 'CEH': why = 'в каталоге PTHK размеров e и H нет'
        elif vals: why = 'в каталоге этой колонки нет'
        elif brand == 'Sant': why = 'в каталогах SANT такого артикула нет'
        elif m: why = 'этого типоразмера (диаметр + длина) в каталоге MetalCraft нет'
        else: why = 'артикула нет в каталогах'
        notdone.append([art, brand, ', '.join(still), why])

wb.save(D + 'Державки_по_каталогу.xlsx')
rep = openpyxl.Workbook()
def sheet(w, title, head, rows, widths):
    w.title = title; w.append(head)
    for c in w[1]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDE6F0'); c.alignment = Alignment(wrap_text=True)
    for rr in rows: w.append(rr)
    for i, wd in enumerate(widths, 1): w.column_dimensions[get_column_letter(i)].width = wd
    for rw in w.iter_rows(min_row=2):
        for c in rw: c.alignment = Alignment(wrap_text=True, vertical='top')
    w.freeze_panes = 'A2'
sheet(rep.active, 'Заполнено (было пусто)', ['Артикул', 'Бренд', 'Что вписано', 'Как получено', 'Источник'], filled, [20, 12, 60, 28, 46])
sheet(rep.create_sheet(), 'Исправлено (было ≠ каталог)', ['Артикул', 'Бренд', 'Было → стало', 'Источник'], fixed, [20, 12, 70, 46])
sheet(rep.create_sheet(), 'Не заполнено', ['Артикул', 'Бренд', 'Пустые колонки', 'Причина'], notdone, [20, 12, 40, 55])
sheet(rep.create_sheet(), 'Оставлено, проверить', ['Артикул', 'Бренд', 'Что не так', 'Источник'], kept, [20, 12, 70, 46])
rep.save(D + 'Отчёт_по_каталогу.xlsx')
nfix = sum(x[2].count('→') for x in fixed)
print('filled rows', len(filled), 'fixed rows', len(fixed), 'fixed cells', nfix, 'notdone', len(notdone), 'kept', len(kept))
