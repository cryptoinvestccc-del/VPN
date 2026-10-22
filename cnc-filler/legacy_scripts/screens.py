import openpyxl,re
from openpyxl.styles import Font, PatternFill, Alignment
B1='S40U-MSKNL12 S40U-MSKNR12 S50U-MSKNL12 S40U-MTFNL16 S50U-MTFNR16 S50U-MTQNR16 S50U-MTUNR16 S40U-MTWNR16 S50U-MTWNR16 S40U-MVQNL16 S40U-MVQNR16 S40T-MVWNL16 S40T-MVWNR16 S40U-MVXNR16'
B2='S20S-MWLNR08 S25R-MWLNR08 S50U-MWLNR08 S50U-MWLNL08 C10M-SCLCR06 E10M-SCLCR06 H10K-SCLCR06 C07K-SCLCR04 H07H-SCLCR04 C08K-SCLCR04 C08K-SCLCR06 E08K-SCLCR04 E08K-SCLCR06 H08J-SCLCR04 C04G-SCLCR03 E04G-SCLCR03 H04F-SCLCR03 C05H-SCLCR03 C05H-SCLCR04 E05H-SCLCR03 H05G-SCLCR03 H05G-SCLCR04 C06J-SCLCR03 C06J-SCLCR04 E06J-SCLCR03 E06J-SCLCR04 E06J-SCLCR06 H06H-SCLCR03 H06H-SCLCR04 S10M-SCLCL06-16 S10M-SCLCR06-16 C12Q-SCLCR06 C16R-SCLCR09'
B3='S16R-SCLCR09H C20S-SCLCR09 S20S-SCLCR09H S25T-SCLCR09H C04G-SDQCR04 E04G-SDQCR04 C05H-SDQCR04 E05H-SDQCR04 C06J-SDQCR04 E06J-SDQCR04 C08K-SDQCR04 C08K-SDQCR07 E08K-SDQCR04 E08K-SDQCR07 C10M-SDQCR07 E10M-SDQCR07 A16Q-SDQCL11 A16Q-SDQCR11 S40T-SDQCL11 S40T-SDQCR11 C04G-SDUCR04 C05H-SDUCR04 C06J-SDUCR04 S07K-SDUCR07 C07K-SDUCR04 A08K-SDUCR07 C08K-SDUCR04 A10K-SDUCR07 S10K-SDUCL07 S10K-SDUCR07'
B4='S20S-SDUCR11 S25U-SDUCR11 S25S-SSKCR12 S32T-STWCL16 S32T-STWCR16 C05H-SVUCR05 C06J-SVUCR05 C07K-SVUCR05 C08K-SVUCR05 C10M-SVUCR05 S12M-SVUCR08 S40U-SVUCR16 S40T-SVWCR16 S12M-SWLCL06 S12M-SWLCR06'
ARTS=(B1+' '+B2+' '+B3+' '+B4).split()
SANT_CODES={'S07K-SDUCR07':'00067133','S10K-SDUCL07':'00057310','S10K-SDUCR07':'00057191'}
COLS=[('Мин. Ø','Мин. диаметр растачивания, мм'),('d','d, мм'),('L','L, мм'),('S','S, мм'),('e','e, мм'),('H','H, мм')]
s=openpyxl.load_workbook('src.xlsx').active; f=openpyxl.load_workbook('Державки_по_каталогу.xlsx').active
up=openpyxl.load_workbook('Загрузка_в_1С.xlsx').active
uh=[c.value for c in up[1]]; U={r[0]:dict(zip(uh,r)) for r in up.iter_rows(min_row=2,values_only=True)}
h=[c.value for c in s[1]]
rows={}
for ra,rb in zip(s.iter_rows(min_row=2,values_only=True),f.iter_rows(min_row=2,values_only=True)):
    da,db=dict(zip(h,ra)),dict(zip(h,rb)); rows.setdefault(da['Артикул'],[]).append((da,db))
NOCAT={'SCLCR09H':'в каталоге MetalCraft SCLCR09H только S20R и S25S','MWLN':'в каталоге MWLN08 только S20R, S25S, S50V',
 'SDUCR07':'в каталоге SDUC07 начинается с S08K','SDUCR11':'в каталоге SDUC11 нет такой длины (S20R, S25S)','SSKCR12':'в каталоге SSKC только под пластину 09',
 'STWC':'в каталоге STWC16 только S16Q, S20R, S25S','SVUCR08':'в каталоге SVU только под пластины 11 и 16','SVUCR16':'в каталоге SVUC16 есть S40T и S50V, S40U нет',
 'SVWCR16':'в каталоге SVWC16 только до S32T','SWLC':'в каталоге SWLC06 начинается с S16Q','-16':'в каталоге исполнение A16 только S06K–S12M (S10K, S12M), S10M нет',
 'MVWN':'в каталоге MVWN16 только до S32T','MVXN':'в каталоге MVXN16 только до S32T','default':'в каталоге MetalCraft такой длины нет (есть S40T/S50V)'}
def reason(art,brand,col):
    if brand=='Sant':
        if art.startswith('S10K-SDUC'): return 'в каталоге SANT (стр. 73) у S08K/S10K колонка e пустая'
        return 'артикула нет в каталогах SANT'
    if art[0] in 'CEH': 
        if art=='C20S-SCLCR09': return 'в PTHK для C20S два варианта Мин. Ø (22 и 25), не выбирал'
        return 'в каталоге PTHK колонок e и H нет'
    if art.startswith('S40T-SDQC'): return 'в базе d = 32, а по артикулу 40 — сначала решить, что верно'
    for k,v in NOCAT.items():
        if k!='default' and k in art: return v
    return NOCAT['default']
out=openpyxl.Workbook(); ws=out.active; ws.title='Пропуски со скринов'
head=['Артикул','Код','Бренд']+[c[0] for c in COLS]+['Что вписал','Что осталось пустым и почему']
ws.append(head)
Y=PatternFill('solid',fgColor='FFF2A8'); G=PatternFill('solid',fgColor='EDEDED')
for art in ARTS:
    for da,db in rows.get(art,[]):
        brand=da['Бренд']; code=da['_Код товара']; code=str(code).zfill(8) if isinstance(code,int) else str(code).strip()
        if art in SANT_CODES and brand=='Sant' and code!=SANT_CODES[art]: continue
        u=U.get(art) if brand!='Sant' else None
        vals=[];filled=[];empty=[]
        for short,c in COLS:
            v=u[c] if (u and c in u) else db[c]
            if v in (None,'','-',0): empty.append(short); vals.append(None)
            else:
                vals.append(v)
                if u and c in u and da[c]!=db[c]: filled.append(f'{short}={v}')
        src=''
        if filled: src=' — из каталога PTHK' if art[0] in 'CEH' else (' — из каталога MetalCraft' if any(x.startswith(('S=','Мин')) for x in filled) else ' — из артикула')
        rs=f"{', '.join(empty)}: {reason(art,brand,None)}" if empty else '—'
        ws.append([art,code,brand]+vals+[(', '.join(filled)+src) if filled else '—', rs])
        r=ws.max_row
        for k,(short,c) in enumerate(COLS):
            cell=ws.cell(r,4+k)
            if vals[k] is None: cell.value='—'; cell.fill=G
            elif f'{short}={vals[k]}' in filled: cell.fill=Y
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDE6F0'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(horizontal='left' if c.column in (1,2,3,10,11) else 'center',vertical='top',wrap_text=c.column>=10)
for i,w in enumerate([18,11,11,9,7,7,7,7,7,30,55],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes='D2'
NOTE={'S20S-MWLNR08':'В Excel есть значения, на сайте пусто. В каталоге S20S нет, источник цифр неизвестен — в загрузку не включал',
'S25R-MWLNR08':'В Excel бренд kkdaoku, на сайте MetalCraft. В каталоге S25R нет, цифры как у Sant — в загрузку не включал',
'S50U-MWLNR08':'В Excel есть значения, на сайте пусто. Цифры один в один как у Sant S50U-MWLNR08 — в загрузку не включал',
'C20S-SCLCR09':'В Excel e 28 / H 18 — не из каталога; Мин. Ø в PTHK 22 или 25 — в загрузку не включал',
'C10M-SCLCR06':'S = 6 из каталога PTHK добавлен в загрузку (на сайте был пуст); e 17 и H 9 — из вашей базы, в каталоге их нет',
'S40T-MVWNL16':'Заполнена в базе, но в каталоге S40T-MVWN нет — источник цифр неизвестен',
'S40U-MTFNL16':'В базе L = 300, а буква U в артикуле = 350; Мин. Ø 50 и H 37 — не из каталога'}
for r in range(2,ws.max_row+1):
    a=ws.cell(r,1).value
    if a in NOTE: ws.cell(r,11).value=(ws.cell(r,11).value if ws.cell(r,11).value!='—' else '')+(' | ' if ws.cell(r,11).value not in ('—',None,'') else '')+NOTE[a]
for art,code in (('S07K-SDUCR07','00067133'),('S10K-SDUCL07','00057310'),('S10K-SDUCR07','00057191')):
    ws.append([art,code,'Sant',13,{'S07K-SDUCR07':7}.get(art,10),125,{'S07K-SDUCR07':8}.get(art,7.7),'—',{'S07K-SDUCR07':6.5}.get(art,9),'—',
      'e: в каталоге SANT S07K нет' if art.startswith('S07K') else 'e: в каталоге SANT (стр. 73) у S08K/S10K колонка e пустая'])
    rr=ws.max_row; ws.cell(rr,8).fill=G
    for c in ws[rr]: c.alignment=Alignment(horizontal='left' if c.column in (1,2,3,10,11) else 'center',vertical='top',wrap_text=c.column>=10)
out.save('Пропуски_со_скринов.xlsx'); print(ws.max_row-1)
