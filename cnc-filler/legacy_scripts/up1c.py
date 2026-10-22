import re, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
SRC='/root/.claude/uploads/77365e1c-a1c8-5008-b397-5c19b94b801a/61939ca6-_____Microsoft_Excel_6.xlsx'
OUT='/tmp/claude-0/-home-claude/77365e1c-a1c8-5008-b397-5c19b94b801a/scratchpad/xl/Пластины_Huareal_загрузка_1С.xlsx'
ws=openpyxl.load_workbook(SRC).active
H=[c.value for c in ws[1]]; N=len(H)
# индексы (0-based)
I=dict(art=0,brand=1,dic=2,iso=3,clr=4,hole=5,S=6,seo=7,mat1=8,code=9,r=10,typ=11,mat2=12,mark=13,oic=14,coat=15,gr=16,dfix=17,shape=18,size=19,L=20,thk=21,d=22,vid=23,rad=24)
# Каталог Huareal 2026-2027: сплав -> (ISO, покрытие, тип пластины), стр. A12-A14
GR={'HR7115':('M S','M, S. (Нержавеющая сталь, Жаропрочные сплавы)','PVD','Твердосплавная'),
    'HR7225':('M','M (Нержавеющая сталь)','PVD','Твердосплавная'),
    'HR8225':('P','P (Сталь)','CVD','Твердосплавная'),
    'HR9205':('S','S (Жаропрочные сплавы)','PVD','Твердосплавная'),
    'HR9215':('S','S (Жаропрочные сплавы)','PVD','Твердосплавная'),
    'HR115C':('P','P (Сталь)','PVD (Металлокерамическое покрытие)','Металлокерамическая')}
GEO={'07':dict(dic='6.35',S=2.38,oic=6.35,dfix='2.8',size=7,L=7.8,thk=2,d=2.8),
     '11':dict(dic='9.525',S=3.97,oic=9.525,dfix='4.4',size=11,L=11.6,thk='T3',d=4.4)}
Y=PatternFill('solid',fgColor='FFF2A8')
out=openpyxl.Workbook(); o=out.active; o.title='Загрузка'; o.append(H)
notes=[]
for row in ws.iter_rows(min_row=2):
    v=[c.value for c in row]
    art=v[0]
    if not art: continue
    code=str(v[9]).strip() if v[9] is not None else ''
    m=re.match(r'^DCMT(07|11)(02|T3)(\d\d)(HQ)?-?(TF|TM)?-(H[RS]\w+)$',art)
    if not m: continue
    sz,th,rr,hq,cb,gr=m.groups()
    if gr=='HR115C' and not hq: continue
    ours = v[1] is None or any(x is None for x in (v[11],v[13],v[15],v[17]))
    if not ours or not code: continue
    new=dict(art=art,brand='Huareal',iso='DCMT',clr='C (7°)',r=int(rr)/10,mark=art.split('-')[0].replace('HQ',''),
             shape='D (ромб 55°)',gr=gr,rad=str(int(rr)/10),**GEO[sz])
    new['vid']='Чистовая' if cb=='TF' or hq else 'Получистовая'
    miss=[]
    if gr in GR:
        a,b,c,t=GR[gr]; new.update(mat1=a,mat2=b,coat=c,typ=t)
    else: miss.append(f'сплав {gr} в каталогах Huareal нет: не заполнены материал, покрытие, тип пластины')
    filled=[]
    vals=list(v)+[None]*(N-len(v))
    for k,val in new.items():
        j=I[k]
        if vals[j] in (None,''):
            vals[j]=val; filled.append(j)
    vals[9]=code.zfill(8)
    o.append(vals); rr_=o.max_row
    o.cell(rr_,10).number_format='@'
    for j in filled: o.cell(rr_,j+1).fill=Y
    src='каталог Huareal 2026-2027: A50 (DCMT…-TF/TM/HQ), A12–A14 (сплавы)'
    if hq and gr=='HR115C': miss.append('HQ в каталоге — «чистовая — получистовая», поставил «Чистовая»; кермет с PVD-покрытием')
    if art=='DCMT11T312-TM-HS6115': miss.append('бренд поставил Huareal, но DCMT11T312-TM-HS7125/HS8125 и DCMT11T304-TM-HS6115 в 1С записаны как Hadsto — проверьте')
    notes.append([code.zfill(8),art,len(filled),'; '.join(miss) or 'всё по каталогу',src])
for c in o[1]: c.font=Font(bold=True); c.alignment=Alignment(wrap_text=True,vertical='center'); c.fill=PatternFill('solid',fgColor='DDE6F0')
o.freeze_panes='B2'
for col,w in zip('ABCDEFGHIJKLMNOPQRSTUVWXY',[24,10,9,8,9,8,6,8,8,10,6,13,30,12,7,8,9,9,13,7,6,7,6,12,8]): o.column_dimensions[col].width=w
n=out.create_sheet('Что заполнено')
n.append(['Код товара','Артикул','Заполнено ячеек','Примечание','Источник'])
for r in notes: n.append(r)
for c in n[1]: c.font=Font(bold=True)
for col,w in zip('ABCDE',[11,24,10,70,50]): n.column_dimensions[col].width=w
for r in n.iter_rows(min_row=2): r[0].number_format='@'
out.save(OUT); print(len(notes))
for r in notes: print(r[:4])
