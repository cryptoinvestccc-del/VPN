import re, sys, openpyxl, copy
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, '/tmp/claude-0/-home-claude/77365e1c-a1c8-5008-b397-5c19b94b801a/scratchpad/xl')
from kk import KK

D = '/tmp/claude-0/-home-claude/77365e1c-a1c8-5008-b397-5c19b94b801a/scratchpad/xl/'
KKP = lambda p: f'каталог MetalCraft, стр. {p} PDF'

# PTHK carbide / HSS small bars: (prefix-series) -> (D, L, A, Dmin, catalog page)
PT = {}
def pt(series, page, rows):
    for pre, d, L, A, dm in rows:
        PT[(pre, series)] = (d, L, A, dm, page)
pt('SCLC03', 9, [('C04G',4,90,4.5,5),('C05H',5,100,5.5,6),('C06J',6,110,6.5,7)])
pt('SCLC04', 9, [('C05H',5,100,5.5,6),('C06J',6,110,6.5,7),('C07K',7,125,7.5,8),('C08K',8,125,8.5,9)])
pt('SCLC06', 9, [('C06J',6,110,7.5,8.5),('C07K',7,125,8,9),('C08K',8,125,9,10),('C08M',8,150,9,10),('C10K',10,125,11,12),('C10M',10,150,11,12),('C12Q',12,180,13,14)])
pt('SCLC03', 10, [('E04G',4,90,4.5,5),('E05H',5,100,5.5,6),('E06J',6,110,6.5,7)])
pt('SCLC04', 10, [('E06J',6,110,6.5,7),('E07K',7,125,7.5,8),('E08K',8,125,8.5,9)])
pt('SCLC06', 10, [('E06J',6,110,6.5,7),('E07K',7,125,7.5,8),('E08M',8,150,9,10),('E08K',8,125,9,10),('E10K',10,125,11,12),('E10M',10,150,11,12),('E12M',12,150,13,14),('E12Q',12,180,13,14)])
pt('SCLC03', 93, [('H04F',4,80,4.5,5),('H05G',5,90,5.5,6),('H06H',6,100,6.5,7)])
pt('SCLC04', 93, [('H05G',5,90,5.5,6),('H06H',6,100,6.5,7),('H07H',7,100,7.5,8),('H08J',8,110,8.5,9)])
pt('SCLC06', 93, [('H07H',7,100,7.7,8.5),('H06H',6,100,7.5,8.5),('H08J',8,110,8.7,9.5),('H10K',10,125,11,12),('H12L',12,140,13,14)])
pt('SCLC06', 9, [('C14Q',14,180,15,16),('C16R',16,200,17,18)])
pt('SCLC09', 9, [('C12Q',12,180,13.5,15),('C14Q',14,180,15,18),('C16Q',16,180,17,18),('C16R',16,200,17,18),('C18R',18,200,19,20),('C18S',18,250,19,20),('C20R',20,200,21,22),('C25S',25,250,28,29),('C25T',25,300,28,29),('C32T',32,300,35,38),('C32U',32,350,35,38)])
pt('SCLC12', 9, [('C32T',32,300,35,38),('C32U',32,350,35,38),('C40U',40,350,44,48),('C40V',40,400,44,48)])
pt('SCLC06', 10, [('E14Q',14,180,15,16),('E16Q',16,180,17,18)])
pt('SCLC09', 10, [('E12Q',12,180,13,15),('E16Q',16,180,17,18),('E16R',16,200,17,18),('E18R',18,200,19,20),('E18S',18,250,19,20),('E20R',20,200,21,22),('E20S',20,250,21,22),('E25T',25,300,27,29)])
pt('SCLC12', 10, [('E32T',32,300,34,38),('E32U',32,350,34,38),('E40U',40,350,44,48),('E40V',40,400,44,48)])
pt('SCLC06', 93, [('H14M',14,150,15,16),('H16Q',16,180,17,18)])
pt('SCLC09', 93, [('H12L',12,140,13.5,15),('H14M',14,150,15,16),('H16Q',16,180,17,18),('H18R',18,200,19,20),('H18S',18,250,19,20),('H20R',20,200,21,22),('H25S',25,250,27,29),('H32T',32,300,34,36)])
pt('SCLC12', 93, [('H40U',40,350,44,48)])
for pre in 'CE':
    pg = {'SDUC': 24, 'SDQC': 25}
    for s in ('SDUC', 'SDQC'):
        pt(s + '04', pg[s], [(pre+'04G',4,90,5.5,6.5),(pre+'05H',5,100,6.5,7.5),(pre+'06J',6,110,7.5,8.5),(pre+'07K',7,125,8.5,9.5),(pre+'08K',8,125,9.5,10.5)])
        pt(s + '07', pg[s], [(pre+'08K',8,125,11,12),(pre+'10M',10,150,13,14),(pre+'14Q',14,180,17,18)])
pt('SVUC05', 35, [('C05H',5,100,8,10),('C06J',6,110,9,10),('C07K',7,125,10,11),('C08K',8,125,11,12),('C10M',10,150,13,14)])

SANT = {  # brand Sant exact rows from SANT catalogue (2024-25 p69 = 2022 p38 PDF)
    'S40T-DCLNL12': ({'e, мм': 50}, 'SANT 2024-25, стр. 69 каталога (PDF 38): S40T-DCLNR/L 12', 'из каталога'),
}

C_D, C_S, C_E, C_DM, C_H, C_L, C_DM2, C_SEC, C_LR = 'd, мм','S, мм','e, мм','Dmin, мм','H, мм','L, мм','Мин. диаметр растачивания, мм','Сечение резца, мм','Длина резца, мм'
CORE = [C_D, C_S, C_E, C_DM, C_H, C_L, C_DM2, C_SEC, C_LR]
E = lambda v: v in (None, '', '-', 0, '0')

def num(v):
    try:
        return float(str(v).replace('Ø', '').replace(',', '.'))
    except Exception:
        return None

def kkey(art):
    m = re.match(r'^([A-Z]\d{2}[A-Z])-([A-Z]{4,5})(\d{2})(H?)(?:-.*)?$', art)
    if not m:
        return None, None
    pre, ser, size, h = m.groups()
    if len(ser) == 5 and ser[-1] in 'RL':
        ser = ser[:-1]
    else:
        return None, None
    return pre, ser, size, h

wb = openpyxl.load_workbook(D + 'src.xlsx')
ws = wb.active
hdr = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(hdr)}
fillY = PatternFill('solid', fgColor='FFF2A8')
fillO = PatternFill('solid', fgColor='FFD8A8')

done, notdone, diffs = [], [], []
for r in range(2, ws.max_row + 1):
    row = {h: ws.cell(r, col[h]).value for h in hdr}
    art = row['Артикул'] or ''
    miss = [k for k in CORE if E(row[k])]
    if not miss:
        continue
    brand = row['Бренд']
    vals, src, how = {}, '', ''
    parts = kkey(art)
    if len(parts) == 4:
        pre, ser, size, h = parts
    else:
        pre = ser = size = h = None
    if art in SANT:
        vals, src, how = SANT[art]
    elif pre and pre[0] == 'S' and brand in ('MetalCraft', 'kkdaoku'):
        key = ser + size + h
        alt = ser[:3] + '_' + size if ser[:2] == 'SV' else None
        hit = KK.get((pre, key)) or (KK.get((pre, alt)) if alt else None)
        if hit:
            (dm, d, L, L1, H, f), page = hit
            vals = {C_D: d, C_S: f, C_E: L1, C_DM: dm, C_H: H, C_L: L, C_DM2: dm, C_SEC: f'Ø{d}', C_LR: L}
            src, how = KKP(page) + f' ({pre}-{ser}R/L{size}{h})', 'из каталога'
    elif pre and pre[0] in 'CEH' and brand == 'MetalCraft':
        hit = PT.get((pre, ser + size))
        if hit:
            d, L, A, dm, page = hit
            s = A - d / 2
            vals = {C_D: d, C_S: int(s) if s == int(s) else s, C_DM: dm, C_L: L, C_DM2: dm, C_SEC: f'Ø{d}', C_LR: L}
            src, how = f'PTHK/Weiding, стр. {page} каталога', 'S = A − d/2 (пересчёт), остальное из каталога'
    if vals and not E(row[C_D]) and num(row[C_D]) != num(vals.get(C_D, row[C_D])):
        diffs.append([art, brand, f'd: на сайте {row[C_D]}, в каталоге {vals[C_D]} — значения не вписывал, сначала проверьте диаметр', src])
        notdone.append([art, brand, ', '.join(k.split(',')[0] for k in CORE if E(row[k])), f'диаметр на сайте ({row[C_D]}) не совпадает с каталогом ({vals[C_D]}), не заполнял'])
        continue
    filled, still, cmp_ = [], [], []
    for k in CORE:
        cur = row[k]
        if k in vals:
            v = vals[k]
            if E(cur):
                c = ws.cell(r, col[k]); c.value = v
                c.fill = fillO if how == 'по аналогу' else fillY
                filled.append(f'{k.split(",")[0]}={v}')
            elif k not in (C_SEC,) and num(cur) is not None and num(v) is not None and abs(num(cur) - num(v)) > 0.01:
                cmp_.append(f'{k.split(",")[0]}: на сайте {cur}, в каталоге {v}')
    # copy Dmin <-> Мин. диаметр within row if one is present
    for a, b in ((C_DM, C_DM2), (C_DM2, C_DM)):
        va, vb = ws.cell(r, col[a]).value, ws.cell(r, col[b]).value
        if E(vb) and not E(va):
            c = ws.cell(r, col[b]); c.value = va; c.fill = fillY
            filled.append(f'{b.split(",")[0]}={va} (из колонки {a.split(",")[0]} этой строки)')
            if not src:
                src, how = 'та же строка таблицы', 'дубль значения Dmin'
    for k in CORE:
        if E(ws.cell(r, col[k]).value):
            still.append(k.split(',')[0])
    if filled:
        done.append([art, brand, ', '.join(filled), how, src])
    if still:
        if not vals:
            why = 'артикула нет ни в одном каталоге'
            if pre and pre[0] == 'S' and brand == 'MetalCraft':
                why = 'серия есть в каталоге MetalCraft, но этого типоразмера (длины/диаметра) там нет' if any(k[1] in (ser + size + h, (ser[:3] + '_' + size) if ser and ser[:2]=='SV' else '') for k in KK) else 'серии нет в каталогах'
            if brand == 'Sant':
                why = 'в каталогах SANT такого артикула нет; у PTHK похожий отличается по размерам'
        elif pre and pre[0] in 'CEH':
            why = 'в каталоге PTHK размеров e и H нет (только чертёж)'
        else:
            why = 'в каталоге этой колонки нет'
        notdone.append([art, brand, ', '.join(still), why])
    if cmp_:
        diffs.append([art, brand, '; '.join(cmp_), src])

wb.save(D + 'out_filled.xlsx')

rep = openpyxl.Workbook()
def sheet(ws_, title, head, rows, widths):
    ws_.title = title
    ws_.append(head)
    for c in ws_[1]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDE6F0')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for rr in rows:
        ws_.append(rr)
    for i, w in enumerate(widths, 1):
        ws_.column_dimensions[get_column_letter(i)].width = w
    for row_ in ws_.iter_rows(min_row=2):
        for c in row_:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws_.freeze_panes = 'A2'
sheet(rep.active, 'Заполнено', ['Артикул', 'Бренд', 'Что вписано', 'Как получено', 'Источник'], done, [20, 12, 60, 34, 46])
sheet(rep.create_sheet(), 'Не заполнено', ['Артикул', 'Бренд', 'Пустые колонки', 'Причина'], notdone, [20, 12, 40, 60])
sheet(rep.create_sheet(), 'Расхождения', ['Артикул', 'Бренд', 'Расхождение (сайт ≠ каталог)', 'Источник'], diffs, [20, 12, 60, 46])
rep.save(D + 'report.xlsx')
print(len(done), len(notdone), len(diffs))
