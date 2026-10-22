import openpyxl,re
FORCE={'C10M-SCLCR06'}
from openpyxl.styles import Font, PatternFill, Alignment
ISO=dict(F=80,G=90,H=100,J=110,K=125,L=140,M=150,N=160,P=170,Q=180,R=200,S=250,T=300,U=350,V=400)
Y=PatternFill('solid',fgColor='FFF2A8')
base=openpyxl.load_workbook('Державки_по_каталогу.xlsx'); bw=base.active
h=[c.value for c in bw[1]]; col={x:i+1 for i,x in enumerate(h)}
for r in range(2,bw.max_row+1):
    a=bw.cell(r,1).value or ''
    m=re.match(r'^S(\d{2})([A-Z])-[A-Z]{4}([RL])\d{2}$',a)
    if not m or bw.cell(r,col['Бренд']).value!='MetalCraft': continue
    d=int(m.group(1)); L=ISO.get(m.group(2))
    want={'d, мм':d,'Сечение резца, мм':f'Ø{d}','Исполнение':'R (правое)' if m.group(3)=='R' else 'L (левое)'}
    if L: want.update({'L, мм':L,'Длина резца, мм':L})
    for k,v in want.items():
        c=bw.cell(r,col[k])
        if c.value in (None,'','-'): c.value=v; c.fill=Y
base.save('Державки_по_каталогу.xlsx')
a=openpyxl.load_workbook('src.xlsx').active; b=openpyxl.load_workbook('Державки_по_каталогу.xlsx').active
hdr=[c.value for c in a[1]]
cols=['Исполнение','d, мм','S, мм','e, мм','Dmin, мм','H, мм','L, мм','Сечение резца, мм','Длина резца, мм','Мин. диаметр растачивания, мм']
idx=[hdr.index(c) for c in cols]; ic=hdr.index('_Код товара')
out=openpyxl.Workbook(); ws=out.active; ws.title='Загрузка в 1С'
ws.append(['Артикул','_Код товара','Бренд']+cols); conf=[]
for ra,rb in zip(a.iter_rows(min_row=2),b.iter_rows(min_row=2)):
    art=rb[0].value or ''
    m=re.match(r'^S(\d{2})([A-Z])-[A-Z]{4}([RL])\d{2}$',art)
    if m and rb[hdr.index('Бренд')].value=='MetalCraft':
        d=str(int(m.group(1))); L=ISO.get(m.group(2))
        for k,v in (('d, мм',d),('Сечение резца, мм',d),('L, мм',L),('Длина резца, мм',L)):
            cur=rb[hdr.index(k)].value; nn=re.search(r'\d+',str(cur)) if cur not in (None,'') else None
            if v and nn and nn.group()!=str(v): conf.append([art,k,cur,v])
        cur=rb[hdr.index('Исполнение')].value
        if cur and str(cur)[0]!=m.group(3): conf.append([art,'Исполнение',cur,m.group(3)])
    if art not in FORCE and all(ra[i].value==rb[i].value for i in range(len(hdr))): continue
    code=rb[ic].value; code=str(code).zfill(8) if isinstance(code,int) else code
    ws.append([art,code,rb[hdr.index('Бренд')].value]+[rb[i].value for i in idx])
    r=ws.max_row; ws.cell(r,2).number_format='@'
    for k,i in enumerate(idx):
        if ra[i].value!=rb[i].value: ws.cell(r,4+k).fill=Y
ws.row_dimensions[1].height=45
for row in ws.iter_rows():
    for c in row: c.alignment=Alignment(horizontal='left' if c.column<=3 else 'center',vertical='center',wrap_text=(c.row==1))
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDE6F0')
for i,w in enumerate([20,12,12,14]+[16]*9,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes='D2'
w2=out.create_sheet('Не совпадает с артикулом'); w2.append(['Артикул','Колонка','В базе','По артикулу'])
for x in conf: w2.append(x)
for i,w in enumerate([20,22,14,14],1): w2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
out.save('Загрузка_в_1С.xlsx'); print(ws.max_row-1,len(conf))
